
import streamlit as st
import datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from io import BytesIO

# -----------------------------------------------------------------------------
# 1. 페이지 및 스타일 설정
# -----------------------------------------------------------------------------
st.set_page_config(page_title="SHIN & KIM | 혁신금융지원 자동화 패키지 v3.0", layout="wide", page_icon="⚖️")

# 스타일링 (세종 아이덴티티)
st.markdown("""
    <style>
    .main-header {font-size: 32px; font-weight: bold; color: #003366; border-bottom: 2px solid #003366; padding-bottom: 10px;}
    .sub-header {font-size: 20px; font-weight: bold; color: #555555;}
    .coach-box {background-color: #f8f9fa; padding: 15px; border-left: 4px solid #003366; margin-bottom: 15px; border-radius: 4px;}
    .vendor-badge {background-color: #e3f2fd; color: #0d47a1; padding: 4px 8px; border-radius: 12px; font-size: 0.8em; font-weight: bold;}
    </style>
""", unsafe_allow_html=True)

# 세션 상태 초기화
if 'app_data' not in st.session_state:
    st.session_state['app_data'] = {
        'company_name': '', 'ceo_name': '', 'biz_num': '', 'service_name': '',
        'selected_vendor': '직접 입력',
        'innovativeness': '', 'consumer_benefit': '',
        'regulation_target': '', 'necessity': '',
        'protection_plan': '', 'security_plan': ''
    }

# -----------------------------------------------------------------------------
# 2. 벤더별 법률 논리 데이터베이스 (Knowledge Base)
# -----------------------------------------------------------------------------
VENDOR_TEMPLATES = {
    "Microsoft 365 Copilot": {
        "service_name": "Microsoft 365 Copilot 기반 사내 업무 자동화 및 AI 비서 서비스",
        "regulation_target": "전자금융감독규정 제15조 제1항 제3호 (망분리) 및 제5호 (비인가 단말기 통제)",
        "innovativeness": (
            "1. [업무 생산성 혁신] LLM(거대언어모델)을 활용하여 메일 작성, 데이터 분석, 보고서 초안 생성 시간을 평균 70% 단축.\n"
            "2. [지능형 검색] 사내 비정형 데이터(규정, 매뉴얼 등)를 자연어로 검색·요약하여 의사결정 속도 획기적 개선."
        ),
        "consumer_benefit": (
            "1. 단순 반복 업무 자동화를 통해 금융상품 상담 등 고부가가치 대고객 서비스에 인력 집중 가능.\n"
            "2. 고객 민원 및 질의에 대한 AI 기반 실시간 분석으로 응대 정확도 및 속도 향상."
        ),
        "necessity": (
            "현행 망분리 규제상 내부망에서의 외부 생성형 AI(OpenAI API 등) 연결이 원천 차단되어 있음. "
            "글로벌 금융 경쟁력 확보를 위해 물리적 망분리의 예외를 인정받아 안전한 SaaS 활용 모델을 실증해야 함."
        ),
        "security_plan": (
            "1. [데이터 유출 방지] Microsoft Tenant Isolation 적용 및 고객 데이터의 AI 학습 목적 재사용 금지 확약(No Training Policy).\n"
            "2. [접근 통제] 내부망에서 전용 보안 게이트웨이(Secure Gateway)를 통해서만 Copilot 서버에 접근 허용.\n"
            "3. [정보 보호] 개인정보 포함 파일 업로드 차단 및 실시간 마스킹 시스템 적용."
        )
    },
    "Salesforce": {
        "service_name": "클라우드 기반 통합 CRM(고객관계관리) 및 초개인화 마케팅 플랫폼",
        "regulation_target": "전자금융감독규정 제14조의2 (클라우드컴퓨팅서비스 이용절차) 및 망분리 규제",
        "innovativeness": (
            "1. [데이터 통합] 분산된 고객 데이터를 클라우드 단일 플랫폼(Single Source of Truth)으로 통합하여 360도 고객 뷰 확보.\n"
            "2. [초개인화] AI(Einstein) 분석을 통해 고객 생애 주기별 맞춤형 금융 상품 추천 및 이탈 예측."
        ),
        "consumer_benefit": (
            "1. 고객 니즈에 부합하는 적시(Right-Time) 금융 상품 제안으로 고객 경험(CX) 개선.\n"
            "2. 옴니채널(앱, 웹, 콜센터) 데이터 연동으로 끊김 없는 상담 서비스 제공."
        ),
        "necessity": (
            "고객 데이터의 실시간 분석과 글로벌 표준 CRM 도입을 위해서는 SaaS 형태의 클라우드 이용이 필수적이나, "
            "개인신용정보의 내부망 처리 원칙과 상충되어 규제 특례가 필요함."
        ),
        "security_plan": (
            "1. [암호화] Salesforce Shield 적용을 통한 DB 및 전송 구간 강력 암호화(AES-256) 및 키 관리(BYOK).\n"
            "2. [국외 이전 통제] 국내 리전(Region) 사용을 원칙으로 하며, 국외 이전 필요 시 금융위 사전 승인 및 동의 절차 준수."
        )
    },
    "Palantir Foundry": {
        "service_name": "빅데이터 통합 분석 플랫폼(Foundry) 기반 차세대 자금세탁방지(AML) 및 이상거래탐지(FDS) 시스템",
        "regulation_target": "신용정보업감독규정 및 전자금융감독규정 (외부주문 등에 관한 기준)",
        "innovativeness": (
            "1. [데이터 사일로 제거] 이기종 시스템 간 데이터를 온톨로지(Ontology) 형태로 통합하여 복잡한 자금 흐름 추적.\n"
            "2. [탐지율 제고] 그래프 분석 기술을 활용하여 기존 Rule-based 시스템이 놓치던 지능형 금융 사기 패턴 탐지."
        ),
        "consumer_benefit": (
            "1. 보이스피싱 등 금융 사기 피해 사전 예방을 통한 소비자 자산 보호 강화.\n"
            "2. 오탐지(False Positive) 감소로 인한 정상 거래 고객의 불편 최소화."
        ),
        "necessity": (
            "고도화된 금융 범죄 대응을 위해 글로벌 선진 분석 솔루션 도입이 시급하나, "
            "해외 솔루션의 특성상 망분리 환경에서의 업데이트 및 유지보수 제약으로 인해 특례 적용이 불가피함."
        ),
        "security_plan": (
            "1. [접근 제어] 목적 기반 접근 통제(Purpose-based Access Control) 및 세분화된 권한 관리.\n"
            "2. [감사 추적] 데이터 조회, 가공, 분석 등 모든 작업 이력에 대한 비가역적 감사 로그 기록 및 모니터링."
        )
    },
    "Celonis": {
        "service_name": "프로세스 마이닝(Process Mining) 기반 내부통제 및 업무 효율화 시스템",
        "regulation_target": "금융회사 지배구조 감독규정 (내부통제기준) 및 망분리 규제",
        "innovativeness": (
            "1. [프로세스 가시화] ERP, 뱅킹 시스템 로그를 분석하여 업무 프로세스를 시각화하고 비효율 구간(Bottle neck) 자동 식별.\n"
            "2. [상시 감시] 횡령, 부정 대출 등 이상 징후를 실시간으로 탐지하여 사고 발생 전 차단."
        ),
        "consumer_benefit": (
            "1. 금융 사고 예방을 통한 금융 시스템 신뢰도 제고 및 고객 피해 방지.\n"
            "2. 업무 처리 속도 개선을 통한 대고객 서비스 리드 타임 단축."
        ),
        "necessity": (
            "실시간 로그 데이터의 클라우드(Celonis EMS) 전송이 필요하나, "
            "내부망의 중요 데이터 외부 전송 제한 규제로 인해 온전한 기능 활용이 불가능함."
        ),
        "security_plan": (
            "1. [데이터 가명화] 클라우드 전송 전, 개인식별정보 및 중요 금융 정보를 프록시 서버에서 가명화/비식별화 처리.\n"
            "2. [전송 통제] 업무 프로세스 분석에 필요한 최소한의 로그 데이터만 화이트리스트 방식으로 전송 허용."
        )
    }
}


# -----------------------------------------------------------------------------
# 3. 함수: 금감원 표준 양식 Word 생성기
# -----------------------------------------------------------------------------
def create_standard_docx(data):
    doc = Document()
    
    # 폰트 설정
    style = doc.styles['Normal']
    style.font.name = 'Malgun Gothic'
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Malgun Gothic')

    # [타이틀]
    title = doc.add_heading('혁신금융서비스 지정 신청서', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f'신청일: {datetime.date.today().strftime("%Y년 %m월 %d일")}\n').alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # [1. 신청인 현황]
    doc.add_heading('1. 신청인 일반 현황', level=1)
    table = doc.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    
    cells = table.rows[0].cells
    cells[0].text = '회 사 명'
    cells[1].text = data.get('company_name', '')
    
    cells = table.rows[1].cells
    cells[0].text = '대 표 자'
    cells[1].text = data.get('ceo_name', '')
    
    cells = table.rows[2].cells
    cells[0].text = '사업자등록번호'
    cells[1].text = data.get('biz_num', '')

    doc.add_paragraph('') # 공백

    # [2. 서비스 개요]
    doc.add_heading('2. 서비스 개요 및 규제특례 대상', level=1)
    doc.add_paragraph(f"■ 대상 솔루션: {data.get('selected_vendor', '')}", style='List Bullet')
    doc.add_paragraph(f"■ 서비스 명칭: {data.get('service_name', '')}", style='List Bullet')
    doc.add_paragraph(f"■ 규제특례 대상 법령:\n{data.get('regulation_target', '')}", style='List Bullet')

    # [3. 서비스의 혁신성]
    doc.add_heading('3. 서비스의 혁신성 및 소비자 편익 (법 제13조제4항제2호·3호)', level=1)
    doc.add_paragraph("가. 기존 서비스 대비 차별성 및 혁신성")
    p = doc.add_paragraph(data.get('innovativeness', ''))
    p.paragraph_format.left_indent = Inches(0.2)
    
    doc.add_paragraph("\n나. 금융소비자 편익 증대 효과")
    p = doc.add_paragraph(data.get('consumer_benefit', ''))
    p.paragraph_format.left_indent = Inches(0.2)

    # [4. 규제특례의 불가피성]
    doc.add_heading('4. 규제특례 적용의 불가피성 (법 제13조제4항제4호)', level=1)
    p = doc.add_paragraph(data.get('necessity', ''))
    p.paragraph_format.left_indent = Inches(0.2)

    # [5. 소비자 보호 및 보안]
    doc.add_heading('5. 소비자 보호 및 위험 관리 방안 (법 제13조제4항제7호)', level=1)
    doc.add_paragraph("가. 소비자 보호 및 위험 고지 계획")
    p = doc.add_paragraph(data.get('protection_plan', ''))
    p.paragraph_format.left_indent = Inches(0.2)
    
    doc.add_paragraph("\n나. 금융보안 및 정보보호 대책 (망분리 예외 등)")
    p = doc.add_paragraph(data.get('security_plan', ''))
    p.paragraph_format.left_indent = Inches(0.2)

    # [서명]
    doc.add_paragraph('\n\n위와 같이 혁신금융서비스 지정을 신청합니다.', style='Normal').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f'\n신청인: {data.get("company_name", "")} 대표이사 {data.get("ceo_name", "")} (인)', style='Normal').alignment = WD_ALIGN_PARAGRAPH.RIGHT

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# -----------------------------------------------------------------------------
# 4. 사이드바 구성
# -----------------------------------------------------------------------------
with st.sidebar:
    st.image("https://www.shinkim.com/assets/images/common/logo_ko.png", width=180)
    st.markdown("### 금융규제그룹 혁신금융 TF")
    
    menu = st.radio("진행 단계", [
        "1. 홈 (Dashboard)", 
        "2. [Step 1] 기본 정보 및 자가진단", 
        "3. [Step 2] 핵심 심사요건 수립", 
        "4. [Step 3] 소비자 보호 및 보안", 
        "5. [Step 4] 신청서 생성 및 다운로드"
    ])
    
    st.divider()
    st.info("""
    **👨‍⚖️ Sejong Expert Tip**
    
    글로벌 SaaS 솔루션 도입은
    **'망분리 규제'**와 **'데이터 국외이전'** 이슈가 핵심입니다.
    세종이 제공하는 벤더별 맞춤 논리를 활용하세요.
    """)

# -----------------------------------------------------------------------------
# 5. 메인 화면 로직
# -----------------------------------------------------------------------------

# --- [1. 홈] ---
if menu == "1. 홈 (Dashboard)":
    st.markdown('<div class="main-header">법무법인 세종 혁신금융지원 자동화 패키지 v3.0</div>', unsafe_allow_html=True)
    st.write("")
    
    st.markdown("""
    <div class="coach-box">
    <strong>👋 환영합니다. 법무법인 세종 금융규제그룹입니다.</strong><br>
    본 솔루션은 <strong>글로벌 SaaS(MS Copilot, Salesforce 등)</strong> 도입을 희망하는 금융회사를 위해 특화되었습니다.<br>
    복잡한 외산 솔루션의 규제 리스크를 해소하고, <strong>[혁신금융서비스 지정 매뉴얼]</strong>에 최적화된 신청서를 완성해 드립니다.
    </div>
    """, unsafe_allow_html=True)

    col1, col2 = st.columns(2)
    with col1:
        st.subheader("🌍 지원 벤더 (Supported Vendors)")
        st.markdown("""
        - **Microsoft 365 Copilot**: 생성형 AI 및 망분리 예외
        - **Salesforce**: 클라우드 CRM 및 개인정보 처리
        - **Palantir**: 빅데이터 분석 및 AML/FDS
        - **Celonis**: 프로세스 마이닝 및 내부통제
        """)
    with col2:
        st.subheader("🏆 세종 Track Record")
        st.markdown("""
        - 글로벌 빅테크 기업 규제 자문 **국내 1위**
        - SaaS 망분리 샌드박스 승인 사례 다수 보유
        - 금융보안원 보안성 심의 **100% 통과** 지원
        """)
    
    if st.button("🚀 신청서 작성 시작하기", type="primary"):
        st.toast("Step 1으로 이동합니다. 사이드바 메뉴를 확인하세요.")

# --- [2. Step 1: 기본 정보 및 자가진단] ---
elif menu == "2. [Step 1] 기본 정보 및 자가진단":
    st.title("Step 1. 기본 정보 및 자가진단")
    
    st.subheader("1. 신청인 기본 정보")
    c1, c2, c3 = st.columns(3)
    
    st.session_state['app_data']['company_name'] = c1.text_input("회사명", value=st.session_state['app_data']['company_name'])
    st.session_state['app_data']['ceo_name'] = c2.text_input("대표자명", value=st.session_state['app_data']['ceo_name'])
    st.session_state['app_data']['biz_num'] = c3.text_input("사업자번호", value=st.session_state['app_data']['biz_num'])

    st.subheader("2. 신청 자격 자가진단")
    st.markdown("""
    <div class="coach-box">
    <b>💡 [금융위 사무관 출신 변호사의 조언]</b><br>
    글로벌 벤더 솔루션을 도입하더라도 신청 주체는 반드시 <b>국내 금융회사 또는 핀테크 기업</b>이어야 합니다.
    해외 본사가 직접 신청하는 것은 불가능함을 유의하십시오.
    </div>
    """, unsafe_allow_html=True)

    check1 = st.checkbox("국내에 영업소를 둔 「상법」상 회사입니까?")
    check2 = st.checkbox("제공하려는 서비스가 '금융업' 또는 이와 관련된 업무입니까?")
    check3 = st.checkbox("현행 금융관련법령(망분리 등)으로 인해 서비스 제공이 불가능합니까?")

    if check1 and check2 and check3:
        st.success("✅ 신청 자격 요건을 충족합니다. 다음 단계로 진행하세요.")
    else:
        st.warning("⚠️ 모든 요건을 충족해야 신청이 가능합니다.")

# --- [3. Step 2: 핵심 심사요건 수립] ---
elif menu == "3. [Step 2] 핵심 심사요건 수립":
    st.title("Step 2. 핵심 심사요건 전략 수립")
    
    # 벤더 선택 섹션 (v3.0 핵심 기능)
    st.subheader("0. 도입 솔루션 선택")
    vendor_list = ["직접 입력"] + list(VENDOR_TEMPLATES.keys())
    selected_vendor = st.selectbox("도입 예정인 글로벌 SaaS 솔루션을 선택하세요.", vendor_list, index=0)
    st.session_state['app_data']['selected_vendor'] = selected_vendor

    # 벤더별 인사이트 제공
    if selected_vendor in VENDOR_TEMPLATES:
        st.markdown(f"""
        <div class="coach-box">
        <b>💡 [Sejong Insight: {selected_vendor} 도입 전략]</b><br>
        {selected_vendor} 도입 시 가장 큰 허들은 <b>데이터 국외 이전</b>과 <b>망분리 예외</b>입니다.
        아래 <b>'✨ 세종 전문가 예시 적용'</b> 버튼을 누르면, 금융당국을 설득할 수 있는 
        최적화된 법률 논리와 문구(Best Practice)가 자동으로 입력됩니다.
        </div>
        """, unsafe_allow_html=True)
    
    # 자동 완성 기능 버튼
    btn_label = f"✨ 세종 전문가 예시 적용 ({selected_vendor} Case)" if selected_vendor != "직접 입력" else "✨ 기본 예시 적용"
    if st.button(btn_label):
        if selected_vendor in VENDOR_TEMPLATES:
            template = VENDOR_TEMPLATES[selected_vendor]
            st.session_state['app_data'].update(template)
        else:
            # 기본 예시
            st.session_state['app_data']['service_name'] = "AI 기반 금융 서비스"
            st.session_state['app_data']['regulation_target'] = "전자금융감독규정 등"
            st.session_state['app_data']['innovativeness'] = "기존 방식 대비 효율성 증대..."
            st.session_state['app_data']['consumer_benefit'] = "소비자 편익 증진..."
            st.session_state['app_data']['necessity'] = "현행 규제상 불가능함..."
        st.rerun()

    st.markdown("---")
    
    st.text_input("서비스 명칭", key='service_name', value=st.session_state['app_data']['service_name'])
    st.text_area("규제특례 대상 법령 (구체적 조항)", key='regulation_target', value=st.session_state['app_data']['regulation_target'])

    st.subheader("1. 서비스의 혁신성 및 차별성")
    st.text_area("기존 서비스 대비 기술적/사업적 혁신성을 기술하세요.", key='innovativeness', height=150, value=st.session_state['app_data']['innovativeness'])
    
    st.subheader("2. 소비자 편익")
    st.text_area("비용 절감, 시간 단축, 접근성 향상 등 구체적 편익을 기술하세요.", key='consumer_benefit', height=100, value=st.session_state['app_data']['consumer_benefit'])

    st.subheader("3. 규제특례의 불가피성")
    st.markdown("""
    <div class="coach-box">
    <b>💡 [IT 규제 전문 변호사의 조언]</b><br>
    SaaS 도입은 <b>'클라우드 이용 절차'</b>와 <b>'물리적 망분리 원칙'</b>의 예외를 인정받아야 합니다.
    보안성을 유지하면서도 왜 반드시 이 솔루션을 써야 하는지, 대체 불가능성을 강조하십시오.
    </div>
    """, unsafe_allow_html=True)
    st.text_area("규제로 인한 사업 수행의 어려움과 특례 필요성", key='necessity', height=150, value=st.session_state['app_data']['necessity'])


# --- [4. Step 3: 소비자 보호 및 보안] ---
elif menu == "4. [Step 3] 소비자 보호 및 보안":
    st.title("Step 3. 소비자 보호 및 보안 대책")

    current_vendor = st.session_state['app_data'].get('selected_vendor', '직접 입력')
    
    # 보안 대책 자동 완성 (벤더별 특화)
    if st.button(f"✨ 보안/보호 대책 예시 적용 ({current_vendor})"):
        if current_vendor in VENDOR_TEMPLATES:
            st.session_state['app_data']['security_plan'] = VENDOR_TEMPLATES[current_vendor]['security_plan']
        
        # 공통 소비자 보호 대책
        st.session_state['app_data']['protection_plan'] = (
            "1. [책임보험 가입] 전문인배상책임보험(보상한도 10억 원 이상) 가입을 통해 사고 발생 시 실질적 배상 능력 확보.\n"
            "2. [오인 방지] 앱/웹 실행 시 '금융위 혁신금융서비스 시범 운영' 문구 팝업 게시 및 이용자 동의 징구.\n"
            "3. [민원 대응] 전담 민원 처리반 운영 및 24시간 내 초동 응대 체계 구축."
        )
        st.rerun()

    st.markdown("---")
    
    st.subheader("1. 소비자 보호 및 위험 관리 방안")
    st.text_area("손해배상, 오인방지, 민원처리 계획 등", key='protection_plan', height=150, value=st.session_state['app_data']['protection_plan'])

    st.subheader("2. 금융보안 및 정보보호 대책")
    st.markdown("""
    <div class="coach-box">
    <b>💡 [CISO 출신 전문위원의 조언]</b><br>
    글로벌 SaaS는 <b>데이터 암호화(BYOK 권장), 접근 통제(MFA), 로그 기록</b>이 필수입니다.
    특히 AI 서비스의 경우, <b>학습 데이터 통제 방안</b>을 반드시 명시해야 합니다.
    </div>
    """, unsafe_allow_html=True)
    st.text_area("해킹 방지, 망분리 대체 통제, 개인정보보호 조치 등", key='security_plan', height=150, value=st.session_state['app_data']['security_plan'])

# --- [5. Step 4: 신청서 생성 및 다운로드] ---
elif menu == "5. [Step 4] 신청서 생성 및 다운로드":
    st.title("Step 4. 최종 신청서 생성 및 확인")
    
    st.markdown("### 📄 작성 내용 미리보기")
    
    with st.container(border=True):
        data = st.session_state['app_data']
        
        if not data.get('company_name'):
            st.warning("⚠️ 작성된 내용이 없습니다. 이전 단계로 돌아가 내용을 입력해주세요.")
        else:
            st.markdown(f"**[신청인]** {data.get('company_name')} (대표: {data.get('ceo_name')})")
            if data.get('selected_vendor') != '직접 입력':
                st.markdown(f"<span class='vendor-badge'>{data.get('selected_vendor')} 도입</span>", unsafe_allow_html=True)
            
            st.markdown(f"**[서비스명]** {data.get('service_name')}")
            st.markdown("---")
            st.markdown("**1. 혁신성 및 차별성**")
            st.info(data.get('innovativeness', '입력되지 않음'))
            st.markdown("**2. 규제특례 불가피성**")
            st.info(data.get('necessity', '입력되지 않음'))
            st.markdown("**3. 보안 대책**")
            st.info(data.get('security_plan', '입력되지 않음'))
    
    st.divider()

    col1, col2 = st.columns([2, 1])
    with col1:
        st.markdown("""
        **[법무법인 세종 혁신금융 TF의 마무리 제언]**
        1. 글로벌 SaaS 도입 시 **'클라우드 안전성 평가(CSP 평가)'** 서류가 필수적으로 요구됩니다.
        2. 생성된 초안은 세종 담당 변호사와 최종 검토 후 금융위에 제출하시기 바랍니다.
        """)
        
    with col2:
        if data.get('company_name'):
            docx_file = create_standard_docx(data)
            st.download_button(
                label="📄 혁신금융서비스 신청서(.docx) 다운로드",
                data=docx_file,
                file_name=f"혁신금융신청서_{data.get('company_name')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary",
                use_container_width=True
            )
        else:
            st.button("📄 신청서 생성 불가 (데이터 누락)", disabled=True, use_container_width=True)

# -----------------------------------------------------------------------------
# Footer
# -----------------------------------------------------------------------------
st.markdown("---")
st.markdown("""
<div style="text-align: center; color: grey; font-size: 12px;">
    <strong>법무법인 세종 (SHIN & KIM LLC)</strong> | 금융규제그룹 혁신금융 TF<br>
    서울시 종로구 종로3길 17 D타워 D1 23층<br>
    Copyright © SHIN & KIM LLC. All rights reserved.
</div>
""", unsafe_allow_html=True)
